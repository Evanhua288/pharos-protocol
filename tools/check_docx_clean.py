# -*- coding: utf-8 -*-
r"""check_docx_clean.py —— 出 docx 后的**三项清洁判据**：
   ① 无字面 \uXXXX 转义  ② 无 `****` 标记泄漏  ③ 落款/基线**不得声明过期全值锚**
   用法：python check_docx_clean.py <目录> [<目录> ...]
   退出码：0 = 全过；1 = 有不合格件。"""
import os
import re
import sys

try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass

from docx import Document

try:
    sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
    import anchor_ledger as _AL
    _AL.load(spec_path=os.path.join(os.path.dirname(os.path.dirname(
        os.path.abspath(__file__))), "spec", "CORE-001-V1.2.md"))
except Exception:
    _AL = None

# ⭐ 现行锚（唯一）——**由发布器自动填**（`release2.py` 会改写本行）
CUR = "821b1e52ef160c5f04d2236b5c62a0ae07e1f30425a03ee4dd28dd37fe2926a2"
# 历史前缀：作“当时值”引用属**有意**（不报）。⏪ V1.63：补齐实测历史锚（此前缺 ae7d95db/4ef5ad6a/7d5a8107）
HIST = ("144211bd", "9fb3b4a2", "e7a7c803", "c621d54a", "7c8fa325", "514b248f", "20e66529",
        "c29381f9", "a5fed421", "b05b2d37", "a501ee55", "9a9cc047",
        "a24bce77", "d41d8cd9", "e3b0c442",
        "ae7d95db", "4ef5ad6a", "7d5a8107", "37e03024", "a34c6fb4", "e29c277a",
        "89682277", "18e0fdbd", "7505f6c5", "fd648525", "fe4f8f38", "a3ab0392", "19a0be17")
# ⏪ V1.63：允许**点名报告污染值**的件（更正说明／登记补全等）——由发布器传 `--allow <子串>`
ALLOW = ()
PAT_ESC = re.compile(r"\\u[0-9a-fA-F]{4}")
PAT_FULL = re.compile(r"\b[0-9a-f]{64}\b")

SELF = os.path.basename(__file__)


def text_of(p):
    d = Document(p)
    t = "".join(x.text for x in d.paragraphs)
    for tb in d.tables:
        for r in tb.rows:
            for c in r.cells:
                t += c.text
    return t


def check(p):
    """\u8fd4\u56de\uff08\u5931\u8d25\u9879\u5217\u8868\uff09\u3002
    \u2b50 \u6536\u7d27\u53e3\u5f84\uff08V1.1\uff09\uff1a
      \u2460 `ESC_LEAK`\uff1a**\u4efb\u4f55**\u5b57\u9762 `\\uXXXX` \u2014\u2014**\u65e0\u4f8b\u5916**\u3002
      \u2461 `ASTERISK_LEAK`\uff1a**\u884c\u9996/\u3010\u540e**\u51fa\u73b0 `****`\uff08**\u6807\u8bb0\u6cc4\u6f0f**\uff09\uff1b
         \u7eaf\u7c97\u4f53\u6e32\u67d3\uff08\u5982 `\u2605 ****\u6807\u9898****`\uff09**\u4e0d\u62a5**\u3002
      \u2462 `STALE_ANCHOR`\uff1a**\u4ec5\u5728\u201c`specHash`/\u57fa\u7ebf/\u951a\u201d\u8bed\u5883\u884c**\u91cc\u7684 64 \u4f4d\u503c\uff0c
         \u4e14**\u6392\u9664**\u5df2\u77e5\u5386\u53f2\u503c\u4e0e\u201c\u5f53\u65f6\u503c\u201d\u6807\u8bb0\u884c\u3002
      \u2b50 \u2463 `INBOUND`\uff1a**\u6765\u4ef6\uff08`WORKBUDDY-*`\uff09\u4e0d\u505a\u951a\u5224\u5b9a**\u2014\u2014\u5b83\u91cc\u9762\u7684\u54c8\u5e0c\u662f\u5bf9\u65b9\u81ea\u5df1\u6587\u4ef6\u7684\u54c8\u5e0c\uff0c
         **\u4e0d\u662f\u6211\u65b9\u7684 `specHash`**\uff1b\u82e5\u5f3a\u884c\u5224\u5b9a\u5fc5\u4ea7\u751f\u8bef\u62a5\u3002**\u4f46 `ESC_LEAK`\u4ecd\u68c0**\u3002
    """
    bad = []
    inbound = os.path.basename(p).upper().startswith("WORKBUDDY-")
    try:
        d = Document(p)
        paras = [x.text for x in d.paragraphs]
        t = "".join(paras)
        for tb in d.tables:
            for r in tb.rows:
                for c in r.cells:
                    t += c.text
    except Exception as e:
        return [("UNREADABLE", str(e)[:60])]
    m = PAT_ESC.findall(t)
    if m:
        bad.append(("ESC_LEAK", "%d \u5904\uff1a\u4f8b %s" % (len(m), m[0])))
    if inbound:
        return bad
    for ln in paras:
        if re.match(r"^\s*\*{4,}", ln) or re.search(r"\u3010\s*\*{4,}", ln):
            bad.append(("ASTERISK_LEAK", ln.strip()[:50]))
            break
    for ln in paras:
        if not re.search(r"specHash|\u57fa\u7ebf|\u951a", ln):
            continue
        if "\u5f53\u65f6\u503c" in ln or "\u5386\u53f2" in ln or "\u5f53\u65f6\u8868\u8ff0" in ln:
            continue
        for h in set(PAT_FULL.findall(ln)):
            if _AL is not None:
                _k = _AL.kind(h)
                _cur = _AL.kind(CUR)
                if _k in ("current", "historical") or h == CUR:
                    continue
                if _k == "illegal":
                    bad.append(("ILLEGAL_VALUE", h[:16] + "\u2026"))
                    continue
            elif h == CUR or any(h.startswith(x) for x in HIST):
                continue
            bad.append(("STALE_ANCHOR", h[:16] + "\u2026"))
    return bad


def main(argv):
    # ⏪ V1.63：解析 `--cur <64hex>` 与 `--allow <子串>`（可多次）
    global CUR, ALLOW
    dirs, allow, i = [], [], 1
    while i < len(argv):
        a = argv[i]
        if a == "--cur" and i + 1 < len(argv):
            CUR = argv[i + 1]
            i += 2
        elif a == "--allow" and i + 1 < len(argv):
            allow.append(argv[i + 1])
            i += 2
        else:
            dirs.append(a)
            i += 1
    ALLOW = tuple(allow)
    dirs = dirs or ["."]
    nf = 0
    bad_files = 0
    for D in dirs:
        if os.path.isfile(D) and D.endswith(".docx"):
            files = [D]
        elif os.path.isdir(D):
            files = [os.path.join(D, f) for f in sorted(os.listdir(D))
                     if f.endswith(".docx") and f != SELF
                     and not any(x in f for x in ALLOW)]
        else:
            print("  [SKIP] %s" % D)
            continue
        for p in files:
            nf += 1
            bad = check(p)
            if bad:
                bad_files += 1
                print("  [FAIL] %s" % os.path.basename(p))
                for k, v in bad:
                    print("         %-16s %s" % (k, v))
    print("[DOCXCLEAN] \u626b %d \u4ef6\uff1b\u4e0d\u5408\u683c %d" % (nf, bad_files))
    print("[RESULT] %s" % ("PASS -- \u5168\u90e8\u6e05\u6d01" if not bad_files else "FAIL -- \u6709\u4e0d\u5408\u683c\u4ef6"))
    return 0 if not bad_files else 1


if __name__ == "__main__":
    sys.exit(main(sys.argv))
